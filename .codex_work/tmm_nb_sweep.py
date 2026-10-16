from pathlib import Path
import cmath

import numpy as np
from openpyxl import load_workbook


SOURCE = Path(r"C:\Users\lunarianmetal\Documents\신진 필름개발팀\01. 설비 균일도, 마그넷 조정\굴절률.xlsx")


def read_index(sheet_name):
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = [(float(a), float(b)) for a, b in ws.iter_rows(min_row=2, max_row=40, max_col=2, values_only=True)]
    wb.close()
    return np.array([r[0] for r in rows]), np.array([r[1] for r in rows])


wl_pet, n_pet_raw = read_index("PET_TOYOBO")
wl_nb, n_nb_raw = read_index("Nb2O5")
wl_sio, n_sio_raw = read_index("SiO2")


def n_at(wavelength, grid, values):
    return float(np.interp(wavelength, grid, values))


def reflectance(wavelength, d_sio, d_nb):
    n0 = 1.0
    n1 = n_at(wavelength, wl_sio, n_sio_raw)
    n2 = n_at(wavelength, wl_nb, n_nb_raw)
    n3 = n_at(wavelength, wl_pet, n_pet_raw)

    r01 = (n0 - n1) / (n0 + n1)
    r12 = (n1 - n2) / (n1 + n2)
    r23 = (n2 - n3) / (n2 + n3)
    delta1 = 2 * np.pi * n1 * d_sio / wavelength
    delta2 = 2 * np.pi * n2 * d_nb / wavelength

    r1_stack = (r12 + r23 * cmath.exp(2j * delta2)) / (
        1 + r12 * r23 * cmath.exp(2j * delta2)
    )
    r_total = (r01 + r1_stack * cmath.exp(2j * delta1)) / (
        1 + r01 * r1_stack * cmath.exp(2j * delta1)
    )
    return abs(r_total) ** 2


def reflectance_matrix(wavelength, d_sio, d_nb):
    n0 = 1.0
    n1 = n_at(wavelength, wl_sio, n_sio_raw)
    n2 = n_at(wavelength, wl_nb, n_nb_raw)
    n3 = n_at(wavelength, wl_pet, n_pet_raw)
    delta1 = 2 * np.pi * n1 * d_sio / wavelength
    delta2 = 2 * np.pi * n2 * d_nb / wavelength

    def layer_matrix(n, delta):
        return np.array(
            [[np.cos(delta), 1j * np.sin(delta) / n],
             [1j * n * np.sin(delta), np.cos(delta)]],
            dtype=complex,
        )

    matrix = layer_matrix(n1, delta1) @ layer_matrix(n2, delta2)
    b = matrix[0, 0] + matrix[0, 1] * n3
    c = matrix[1, 0] + matrix[1, 1] * n3
    r_total = (n0 * b - c) / (n0 * b + c)
    return abs(r_total) ** 2


def minimum_for(d_sio, d_nb, low=430.0, high=550.0, step=0.1):
    wavelengths = np.arange(low, high + step / 2, step)
    values = np.array([reflectance(w, d_sio, d_nb) for w in wavelengths])
    idx = int(np.argmin(values))
    return wavelengths[idx], values[idx]


def invert_stationary_wavelength(target_wavelength, d_nb):
    n1 = n_at(target_wavelength, wl_sio, n_sio_raw)
    phase_guess = target_wavelength / (4 * n1) * 1.18
    lo = max(20.0, phase_guess - 15.0)
    hi = phase_guess + 15.0

    def slope(d_sio):
        h = 0.02
        return (
            reflectance(target_wavelength + h, d_sio, d_nb)
            - reflectance(target_wavelength - h, d_sio, d_nb)
        ) / (2 * h)

    xs = np.linspace(lo, hi, 301)
    brackets = []
    last_x, last_y = xs[0], slope(xs[0])
    for x in xs[1:]:
        y = slope(x)
        if last_y == 0 or y == 0 or last_y * y < 0:
            brackets.append((last_x, x))
        last_x, last_y = x, y
    candidates = []
    for a, b in brackets:
        fa = slope(a)
        for _ in range(60):
            m = (a + b) / 2
            fm = slope(m)
            if fa * fm <= 0:
                b = m
            else:
                a, fa = m, fm
        d = (a + b) / 2
        h = 0.1
        candidates.append(d)
    return min(candidates, key=lambda d: abs(d - phase_guess))


print("d_SiO2_nm,d_Nb2O5_nm,lambda_min_nm,R_min")
for d_nb in range(8, 16):
    wavelength, rmin = minimum_for(95.0, float(d_nb))
    print(f"95.0,{d_nb},{wavelength:.1f},{rmin:.8f}")

print("\nSiO2 sensitivity at Nb2O5=10 nm")
for d_sio in (90.0, 92.5, 95.0, 97.5, 100.0):
    wavelength, rmin = minimum_for(d_sio, 10.0)
    print(f"{d_sio:.1f},10,{wavelength:.1f},{rmin:.8f}")

max_difference = max(
    abs(reflectance(w, 95.0, 10.0) - reflectance_matrix(w, 95.0, 10.0))
    for w in np.arange(430.0, 550.1, 0.1)
)
print(f"\nrecursive_vs_matrix_max_abs_difference={max_difference:.3e}")
for target in (470.0, 473.4, 475.0, 480.0):
    d = invert_stationary_wavelength(target, 10.0)
    actual_w, _ = minimum_for(d, 10.0, target - 20, target + 20, 0.02)
    print(f"inverse target={target:.1f}, d_SiO2={d:.4f}, verified_min={actual_w:.2f}")
