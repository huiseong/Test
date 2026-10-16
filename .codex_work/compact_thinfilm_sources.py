from pathlib import Path
import openpyxl

main_path = Path(r"C:\Users\lunarianmetal\Documents\신진 필름개발팀\01. 설비 균일도, 마그넷 조정\2호기 균일도\20230919_RTR#2_Nb2O5, SiO2 균일도\20230919_RTR#2_Nb2O5, SiO2 균일도.xlsx")
ri_path = Path(r"C:\Users\lunarianmetal\Documents\신진 필름개발팀\01. 설비 균일도, 마그넷 조정\굴절률.xlsx")


def short(v):
    if v is None:
        return ""
    s = str(v).replace("\n", " ")
    return s[:80]


for path in [main_path, ri_path]:
    print(f"BOOK\t{path.name}")
    wb = openpyxl.load_workbook(path, data_only=False, read_only=False)
    for ws in wb.worksheets:
        print(f"SHEET\t{ws.title}\t{ws.max_row}x{ws.max_column}\tcharts={len(ws._charts)}")
        if path == main_path:
            for r in range(1, ws.max_row + 1):
                populated = []
                for c in range(1, ws.max_column + 1):
                    v = ws.cell(r, c).value
                    if v is not None:
                        populated.append(f"{ws.cell(r,c).coordinate}={short(v)}")
                if populated:
                    print("ROW\t" + "\t".join(populated))
        else:
            for r in range(1, min(ws.max_row, 8) + 1):
                print("ROW\t" + "\t".join(f"{ws.cell(r,c).coordinate}={short(ws.cell(r,c).value)}" for c in range(1, ws.max_column + 1)))
            print("TAIL\t" + "\t".join(f"{ws.cell(ws.max_row,c).coordinate}={short(ws.cell(ws.max_row,c).value)}" for c in range(1, ws.max_column + 1)))
