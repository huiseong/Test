from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import openpyxl


FILES = [
    Path(r"C:\Users\lunarianmetal\Documents\신진 필름개발팀\01. 설비 균일도, 마그넷 조정\2호기 균일도\20230919_RTR#2_Nb2O5, SiO2 균일도\20230919_RTR#2_Nb2O5, SiO2 균일도.xlsx"),
    Path(r"C:\Users\lunarianmetal\Documents\신진 필름개발팀\01. 설비 균일도, 마그넷 조정\굴절률.xlsx"),
]


def serialise(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def inspect_book(path: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path, data_only=False, read_only=False)
    cached = openpyxl.load_workbook(path, data_only=True, read_only=False)
    result: dict[str, Any] = {
        "path": str(path),
        "size": path.stat().st_size,
        "sheets": [],
        "defined_names": [str(x) for x in wb.defined_names.values()],
    }
    for ws in wb.worksheets:
        cws = cached[ws.title]
        nonempty = []
        formulas = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    item = {
                        "cell": cell.coordinate,
                        "value": serialise(cell.value),
                        "cached": serialise(cws[cell.coordinate].value),
                        "number_format": cell.number_format,
                    }
                    nonempty.append(item)
                    if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                        formulas.append(item)
        charts = []
        for chart in ws._charts:
            charts.append({
                "type": type(chart).__name__,
                "title": serialise(getattr(chart, "title", None)),
                "series_count": len(getattr(chart, "ser", [])),
                "anchor": serialise(getattr(chart, "anchor", None)),
            })
        result["sheets"].append({
            "title": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
            "freeze_panes": serialise(ws.freeze_panes),
            "tables": list(ws.tables.keys()),
            "charts": charts,
            "nonempty_count": len(nonempty),
            "formula_count": len(formulas),
            "nonempty_preview": nonempty[:350],
            "formulas": formulas[:500],
        })
    return result


if __name__ == "__main__":
    out = [inspect_book(path) for path in FILES]
    print(json.dumps(out, ensure_ascii=False, indent=2))
